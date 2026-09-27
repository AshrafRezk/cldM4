"""Extract and render documents. A long extract is a preview plus an artifact URL."""

from __future__ import annotations

import base64
import io
from typing import Any

from app.errors import CloudiatorError

NAME = "extract_document"
CAPABILITY = "tools.docs"
TIMEOUT_SECONDS = 30
GPU = False

PREVIEW_CHARS = 8_000
_KINDS = {"pdf", "docx", "xlsx", "txt"}


async def run(arguments: dict[str, Any], *, settings, artifacts) -> dict[str, Any]:
    action = arguments.get("action") or "extract"
    if action == "render_pdf":
        return _render_pdf(arguments, artifacts)
    if action == "render_xlsx":
        return _render_xlsx(arguments, artifacts)
    if action != "extract":
        raise CloudiatorError(
            400,
            "invalid_request_error",
            "action must be extract, render_pdf, or render_xlsx.",
            param="action",
        )
    data = _decode(arguments.get("file_base64"), limit=settings.vision_max_bytes)
    kind = _kind(arguments.get("filename") or arguments.get("kind") or "")
    text = _extract(data, kind)
    return _preview(text, artifacts, kind)


def _decode(value: Any, *, limit: int) -> bytes:
    if not isinstance(value, str) or not value:
        raise CloudiatorError(400, "invalid_request_error", "file_base64 is required.", param="file_base64")
    try:
        data = base64.b64decode(value, validate=True)
    except Exception as exc:  # noqa: BLE001
        raise CloudiatorError(400, "invalid_request_error", "file_base64 is not valid.", param="file_base64") from exc
    if len(data) > limit:
        raise CloudiatorError(413, "upload_too_large", f"Upload exceeds {limit} bytes.")
    return data


def _kind(name: str) -> str:
    token = name.lower().rsplit(".", 1)[-1]
    if token not in _KINDS:
        raise CloudiatorError(
            400, "invalid_request_error", "filename must end in pdf, docx, xlsx, or txt.", param="filename"
        )
    return token


def _extract(data: bytes, kind: str) -> str:
    if kind == "pdf":
        return _pdf_text(data)
    if kind == "docx":
        return _docx_text(data)
    if kind == "xlsx":
        return _xlsx_text(data)
    return data.decode("utf-8", errors="replace")


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise CloudiatorError(400, "invalid_request_error", "The PDF could not be read.", param="file_base64") from exc
    parts = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(part for part in parts if part).strip()


def _docx_text(data: bytes) -> str:
    from docx import Document

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise CloudiatorError(400, "invalid_request_error", "The document could not be read.", param="file_base64") from exc
    return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text).strip()


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise CloudiatorError(400, "invalid_request_error", "The spreadsheet could not be read.", param="file_base64") from exc
    lines = []
    for sheet in book.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if cell is None else str(cell) for cell in row]
            if any(cells):
                lines.append("\t".join(cells))
    book.close()
    return "\n".join(lines)


def _preview(text: str, artifacts, kind: str) -> dict[str, Any]:
    truncated = len(text) > PREVIEW_CHARS
    payload = {
        "kind": kind,
        "text": text[:PREVIEW_CHARS],
        "chars": len(text),
        "truncated": truncated,
    }
    if truncated:
        artifact_id = artifacts.save(text.encode("utf-8"), ".txt")
        payload["id"] = artifact_id
        payload["url"] = artifacts.sign(artifact_id)
    return payload


def _render_pdf(arguments: dict[str, Any], artifacts) -> dict[str, Any]:
    text = arguments.get("text")
    if not isinstance(text, str) or not text:
        raise CloudiatorError(400, "invalid_request_error", "text is required.", param="text")
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=14)
    pdf.multi_cell(0, 8, text[:20_000])
    raw = bytes(pdf.output())
    artifact_id = artifacts.save(raw, ".pdf")
    return {"id": artifact_id, "url": artifacts.sign(artifact_id), "kind": "pdf"}


def _render_xlsx(arguments: dict[str, Any], artifacts) -> dict[str, Any]:
    columns = arguments.get("columns")
    rows = arguments.get("rows")
    if not isinstance(columns, list) or not all(isinstance(item, str) for item in columns):
        raise CloudiatorError(400, "invalid_request_error", "columns must be an array of strings.", param="columns")
    if not isinstance(rows, list):
        raise CloudiatorError(400, "invalid_request_error", "rows must be an array.", param="rows")
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.append(columns)
    for row in rows:
        if not isinstance(row, list) or len(row) != len(columns):
            raise CloudiatorError(400, "invalid_request_error", "each row must match columns.", param="rows")
        sheet.append([_cell(value) for value in row])
    buffer = io.BytesIO()
    book.save(buffer)
    artifact_id = artifacts.save(buffer.getvalue(), ".xlsx")
    return {"id": artifact_id, "url": artifacts.sign(artifact_id), "kind": "xlsx"}


def _cell(value: Any) -> str | int | float:
    if isinstance(value, bool) or value is None:
        return "" if value is None else str(value).lower()
    if isinstance(value, (int, float)):
        return value
    return str(value)
